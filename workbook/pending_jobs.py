"""Pending-jobs workbook operations."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from workbook.job_identity import canonical_job_url
from workbook.storage import with_workbook

PENDING_JOB_COLUMNS = ("url", "fetched_at")


@with_workbook()
def _pending_job_records(path: Path) -> list[dict[str, object]]:
    """Read current pending URLs from the two-column workbook."""
    if not path.is_file():
        return []

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = {
            str(value).strip(): index
            for index, value in enumerate(header_row)
            if value not in (None, "")
        }
        url_indexes = [
            index
            for index in (headers.get("url"), headers.get("link"))
            if index is not None
        ]
        fetched_at_index = headers.get("fetched_at")
        if not url_indexes:
            return []

        records = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            url = next(
                (
                    str(row[index]).strip()
                    for index in url_indexes
                    if index < len(row) and row[index] is not None and str(row[index]).strip()
                ),
                "",
            )
            if not url:
                continue
            fetched_at = (
                None
                if fetched_at_index is None or fetched_at_index >= len(row)
                else row[fetched_at_index]
            )
            records.append({"url": url, "fetched_at": fetched_at})
        return records
    finally:
        workbook.close()


@with_workbook(write=True)
def _write_pending_job_records(path: Path, records) -> None:
    """Write the two-column pending-jobs workbook from normalized records."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pending Jobs"
    sheet.append(PENDING_JOB_COLUMNS)
    sheet.column_dimensions["A"].width = 72
    sheet.column_dimensions["B"].width = 22

    for record in records:
        sheet.append([record["url"], record.get("fetched_at")])
        sheet.cell(sheet.max_row, 2).number_format = "yyyy-mm-dd hh:mm:ss"

    workbook.save(path)


@with_workbook(write=True)
def remove_pending_job_url(path: Path, url: str) -> int:
    """Remove every pending row whose URL matches the supplied job URL."""
    target = canonical_job_url(url)
    if not target or not path.is_file():
        return 0

    records = _pending_job_records(path)
    remaining_records = [
        record
        for record in records
        if canonical_job_url(str(record["url"])) != target
    ]
    removed = len(records) - len(remaining_records)
    if removed:
        _write_pending_job_records(path, remaining_records)
    return removed
