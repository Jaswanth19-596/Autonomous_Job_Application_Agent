"""Persist failed jobs before removing them from the pending queue."""

from pathlib import Path

from openpyxl import Workbook, load_workbook

from jobs.jobright import remove_pending_job_url
from workbook.job_identity import canonical_job_id, canonical_job_url


def move_failed_job(pending_path: Path, invalid_path: Path, url: str, reason: str) -> None:
    """Upsert a failure by listing URL; retries safely finish pending removal."""
    canonical = canonical_job_url(url)
    if not canonical:
        raise ValueError("Cannot record a rejected job without a valid job URL.")
    job_id = canonical_job_id({"link": canonical})
    book = load_workbook(invalid_path) if invalid_path.exists() else Workbook()
    try:
        sheet = book.worksheets[0]
        columns = {cell.value: cell.column for cell in sheet[1] if cell.value}
        for name in ("job_id", "url", "reason"):
            if name not in columns:
                columns[name] = max(columns.values(), default=0) + 1
                sheet.cell(1, columns[name], name)
        matches = [
            row for row in range(2, sheet.max_row + 1)
            if sheet.cell(row, columns["job_id"]).value == job_id
            or canonical_job_url(sheet.cell(row, columns["url"]).value) == canonical
        ]
        row = matches[0] if matches else sheet.max_row + 1
        for name, value in {"job_id": job_id, "url": url, "reason": reason}.items():
            sheet.cell(row, columns[name], value)
        for duplicate in reversed(matches[1:]):
            sheet.delete_rows(duplicate)
        invalid_path.parent.mkdir(parents=True, exist_ok=True)
        book.save(invalid_path)
    finally:
        book.close()

    # Do not remove the only copy if saving the failure raises an exception.
    remove_pending_job_url(pending_path, url)
