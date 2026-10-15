"""One-time, resumable copying of the job workbooks into Google Drive."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from workbook.drive import (
    DriveClient, DriveStorageError, PROJECT_ROOT, config_path,
    get_credentials, write_private_json,
)
from workbook.storage import WORKBOOK_NAMES


def _validate_workbook(data: bytes) -> None:
    book = load_workbook(BytesIO(data), read_only=True)
    try:
        if not book.sheetnames:
            raise DriveStorageError("A job workbook has no worksheets.")
    finally:
        book.close()


def migrate(client: DriveClient, *, new_folder: bool = False) -> dict:
    """Copy existing files without deleting local data or overwriting Drive edits."""
    manifest = config_path()
    if manifest.exists():
        config = json.loads(manifest.read_text(encoding="utf-8"))
        # A repeated setup verifies Drive; it never re-uploads stale local files.
        with TemporaryDirectory(prefix="job-agent-verify-") as directory:
            for name in WORKBOOK_NAMES:
                destination = Path(directory) / name
                client.download(config["files"][name], destination)
                _validate_workbook(destination.read_bytes())
        return config

    snapshots = {}
    for name in WORKBOOK_NAMES:
        source = PROJECT_ROOT / "data" / name
        if not source.is_file():
            raise DriveStorageError(f"{name} is missing. Restore the local workbook before migrating.")
        snapshots[name] = source.read_bytes()
        _validate_workbook(snapshots[name])

    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    backup = PROJECT_ROOT / "data/drive-migration-backups" / stamp
    backup.mkdir(parents=True)
    for name, data in snapshots.items():
        (backup / name).write_bytes(data)

    folders = [] if new_folder else client.find(
        "trashed = false and mimeType = 'application/vnd.google-apps.folder' "
        "and appProperties has { key='jobAgentStorage' and value='workbooks-v1' }"
    )
    if len(folders) > 1:
        raise DriveStorageError("Multiple agent Drive folders exist. Use --new-folder to copy into a new private folder.")
    folder = folders[0] if folders else client.create_folder("JobApplicationAgent")
    folder_id = folder["id"]
    files = {}
    with TemporaryDirectory(prefix="job-agent-migrate-") as directory:
        for name in WORKBOOK_NAMES:
            existing = client.find(f"trashed = false and '{folder_id}' in parents and name = '{name}'")
            if len(existing) > 1:
                raise DriveStorageError(f"Multiple Drive files named {name}; refusing to select an arbitrary workbook.")
            if existing:
                file_id = existing[0]["id"]
                destination = Path(directory) / name
                client.download(file_id, destination)
                if destination.read_bytes() != snapshots[name]:
                    raise DriveStorageError(
                        f"Drive already contains different data in {name}. Nothing was overwritten. "
                        "To migrate current local files into a separate folder, stop the agent and rerun with --new-folder."
                    )
            else:
                file_id = client.create_workbook(backup / name, folder_id)["id"]
                destination = Path(directory) / name
                client.download(file_id, destination)
                if hashlib.sha256(destination.read_bytes()).digest() != hashlib.sha256(snapshots[name]).digest():
                    raise DriveStorageError(f"Verification failed for {name}. Drive storage has not been activated.")
            files[name] = file_id

    for name, data in snapshots.items():
        if (PROJECT_ROOT / "data" / name).read_bytes() != data:
            raise DriveStorageError(
                "Local job files changed during migration. Drive storage has not been activated. "
                "Stop the running agent and rerun setup with --new-folder."
            )
    config = {"folder_id": folder_id, "files": files, "migrated_at": stamp}
    # Activate only after all files were copied and verified successfully.
    write_private_json(manifest, config)
    return config


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--new-folder", action="store_true", help="Use a new folder when recovering an incomplete migration.")
    parser.add_argument("--status", action="store_true", help="Verify an existing configuration without interactive sign-in.")
    args = parser.parse_args()
    if args.status and not config_path().is_file():
        parser.exit(1, "Drive storage is not configured yet. Run: uv run python -m workbook.drive_setup\n")
    client = None
    try:
        credentials = get_credentials(interactive=not args.status)
        client = DriveClient(credentials)
        config = migrate(client, new_folder=args.new_folder)
        print("Google Drive storage is ready. Restart the agent before the next batch.")
        print(f"Folder: https://drive.google.com/drive/folders/{config['folder_id']}")
        for name, file_id in config["files"].items():
            print(f"{name}: https://drive.google.com/file/d/{file_id}/view")
    except (DriveStorageError, OSError, ValueError, KeyError) as exc:
        parser.exit(1, f"Drive setup stopped: {exc}\n")
    finally:
        if client is not None:
            client.close()


if __name__ == "__main__":
    main()
