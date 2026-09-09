"""Persist Jobright recommendation URLs in the pending-jobs workbook."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Mapping
from urllib.parse import urljoin, urlsplit

from openpyxl import Workbook, load_workbook

from workbook.job_identity import canonical_job_url


JOBRIGHT_RECOMMENDED_URL = "https://jobright.com/jobs/recommended"
PENDING_JOB_COLUMNS = ("url", "fetched_at")


@dataclass(frozen=True)
class PendingJobUpsertResult:
    """Summary of a Jobright URL write to the pending-jobs workbook."""

    inserted: int = 0
    updated: int = 0


def _is_jobright_url(url: str) -> bool:
    """Return whether a URL belongs to Jobright."""
    hostname = (urlsplit(url).hostname or "").lower().removeprefix("www.")
    return (
        hostname in {"jobright.com", "jobright.ai"}
        or hostname.endswith(".jobright.com")
        or hostname.endswith(".jobright.ai")
    )


def _pending_job_records(path: Path) -> list[dict[str, object]]:
    """Read current pending URLs, including the legacy URL-only import shape."""
    if not path.is_file():
        return []

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = {
            str(value).strip(): index
            for index, value in enumerate(header_row)
            if value not in (None, "")
        }
        url_indexes = [
            index
            for index in (headers.get("url"), headers.get("link"))
            if index is not None
        ]
        fetched_at_index = headers.get("fetched_at")
        if not url_indexes:
            return []

        records = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            url = next(
                (
                    str(row[index]).strip()
                    for index in url_indexes
                    if index < len(row) and row[index] is not None and str(row[index]).strip()
                ),
                "",
            )
            if not url:
                continue
            fetched_at = (
                None
                if fetched_at_index is None or fetched_at_index >= len(row)
                else row[fetched_at_index]
            )
            records.append({"url": url, "fetched_at": fetched_at})
        return records
    finally:
        workbook.close()


def _write_pending_job_records(path: Path, records: Iterable[Mapping[str, object]]) -> None:
    """Write the two-column pending-jobs workbook from normalized records."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pending Jobs"
    sheet.append(PENDING_JOB_COLUMNS)
    sheet.column_dimensions["A"].width = 72
    sheet.column_dimensions["B"].width = 22

    for record in records:
        sheet.append([record["url"], record.get("fetched_at")])
        sheet.cell(sheet.max_row, 2).number_format = "yyyy-mm-dd hh:mm:ss"

    workbook.save(path)


def _jobright_url(job: Mapping[str, Any] | Any, position: int) -> str:
    record = job.model_dump() if hasattr(job, "model_dump") else dict(job)
    url = str(record.get("link") or "").strip()
    if url.startswith("/"):
        url = urljoin(JOBRIGHT_RECOMMENDED_URL, url)
    if not url:
        raise ValueError(f"Jobright result #{position} has no detail URL.")
    if not _is_jobright_url(url):
        raise ValueError(f"Jobright result #{position} must be a Jobright detail URL.")
    return url


def recent_workbook_jobright_urls(path: Path, *, limit: int = 20) -> list[str]:
    """Return recently fetched Jobright detail URLs for collection deduplication."""
    if limit < 1:
        return []

    urls: list[str] = []
    seen: set[str] = set()
    for record in reversed(_pending_job_records(path)):
        url = str(record["url"])
        canonical = canonical_job_url(url)
        if not _is_jobright_url(url) or not canonical or canonical in seen:
            continue
        seen.add(canonical)
        urls.append(url)
        if len(urls) == limit:
            break
    return urls


def format_recent_workbook_jobright_urls_for_prompt(urls: Iterable[str]) -> str:
    """Render saved Jobright URLs as clearly delimited reference data."""
    lines = [f"- {url}" for url in urls if str(url).strip()]
    return "\n".join(lines) if lines else "- No prior Jobright URLs are recorded in the workbook."


def normalize_pending_jobs_workbook(path: Path) -> int:
    """Migrate a pending workbook to the required two-column schema."""
    records = _pending_job_records(path)
    _write_pending_job_records(path, records)
    return len(records)


def persist_jobright_jobs(
    path: Path,
    jobs: Iterable[Mapping[str, Any] | Any],
) -> PendingJobUpsertResult:
    """Save Jobright detail URLs to the two-column pending-jobs workbook."""
    records = _pending_job_records(path)
    record_indexes = {
        canonical_job_url(str(record["url"])): index
        for index, record in enumerate(records)
        if canonical_job_url(str(record["url"]))
    }
    fetched_at = datetime.now().replace(microsecond=0)
    inserted = updated = 0

    for position, job in enumerate(jobs, start=1):
        url = _jobright_url(job, position)
        canonical = canonical_job_url(url)
        existing_index = record_indexes.get(canonical)
        if existing_index is None:
            record_indexes[canonical] = len(records)
            records.append({"url": url, "fetched_at": fetched_at})
            inserted += 1
        else:
            records[existing_index] = {"url": url, "fetched_at": fetched_at}
            updated += 1

    if inserted or updated:
        _write_pending_job_records(path, records)
    return PendingJobUpsertResult(inserted=inserted, updated=updated)


def remove_pending_job_url(path: Path, url: str) -> int:
    """Remove every pending row whose URL matches the supplied job URL."""
    target = canonical_job_url(url)
    if not target or not path.is_file():
        return 0

    records = _pending_job_records(path)
    remaining_records = [
        record
        for record in records
        if canonical_job_url(str(record["url"])) != target
    ]
    removed = len(records) - len(remaining_records)
    if removed:
        _write_pending_job_records(path, remaining_records)
    return removed
