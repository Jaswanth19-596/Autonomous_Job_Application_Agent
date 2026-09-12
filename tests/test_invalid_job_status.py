import ast
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, Mock, patch

from openpyxl import Workbook, load_workbook
from workbook.invalid_jobs import move_failed_job
from workbook.job_identity import canonical_job_id


class InvalidJobStatusTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.pending = Path(self.directory.name) / 'pending_jobs.xlsx'
        self.invalid = self.pending.with_name('invalid_jobs.xlsx')
        self.url = 'https://jobright.ai/jobs/info/6aa3b8415c11cce360365dea'
        self.other = 'https://jobright.ai/jobs/info/another-job'
        self.write(self.pending, [['url', 'fetched_at'], [self.url, 'today'], [self.other, 'yesterday']])

    def write(self, path, rows):
        book = Workbook()
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()

    def rows(self, path):
        book = load_workbook(path)
        try:
            rows = list(book.active.values)
            return [dict(zip(rows[0], row)) for row in rows[1:]]
        finally:
            book.close()

    def load_node(self):
        source = Path(__file__).resolve().parents[1] / 'nodes.py'
        node = next(n for n in ast.parse(source.read_text()).body
                    if isinstance(n, ast.AsyncFunctionDef) and n.name == 'update_job_status')
        node.decorator_list = []
        tool = Mock(name='status_tool')
        tool.ainvoke = AsyncMock(return_value={})
        async def progress(_node, _message, operation, **kwargs):
            return await operation
        namespace = dict(Path=Path, __file__=str(source), move_failed_job=move_failed_job,
                         log_node_stage=Mock(), update_job_status_tool=tool,
                         log_tool_calls=Mock(), log_tool_result=Mock(), await_with_progress=progress)
        exec(compile(ast.Module(body=[node], type_ignores=[]), str(source), 'exec'), namespace)
        return namespace['update_job_status'], tool

    async def test_rejection_moves_current_url_and_updates_legacy_row(self):
        self.write(self.invalid, [['job_id', 'reason'], ['keep', 'Keep reason'],
                   [canonical_job_id({'link': self.url}), 'Old reason']])
        node, tool = self.load_node()
        state = dict(validation_phase='NOT_ELIGIBLE', job_url=self.url,
                     pending_jobs_path=str(self.pending), job_details={'id': 'previous-job'},
                     final_status_message='Closed')
        result = await node(state)
        await node(state)  # A retry must not duplicate the failure.
        self.assertEqual(result['completed_jobs'], [self.url])
        self.assertEqual(self.rows(self.pending), [{'url': self.other, 'fetched_at': 'yesterday'}])
        self.assertEqual(self.rows(self.invalid), [
            {'job_id': 'keep', 'reason': 'Keep reason', 'url': None},
            {'job_id': canonical_job_id({'link': self.url}), 'reason': 'Closed', 'url': self.url}])
        tool.ainvoke.assert_not_called()

    async def test_failed_application_moves_but_success_does_not(self):
        for status in ('Applied', 'Not Applied'):
            node, tool = self.load_node()
            await node(dict(validation_phase='ELIGIBLE', job_url=self.url,
                            pending_jobs_path=str(self.pending), job_details={'id': 'stored-id'},
                            application_status=status, final_status_message='Submission failed'))
            tool.ainvoke.assert_awaited_once()
            if status == 'Applied':
                self.assertFalse(self.invalid.exists())
                self.assertEqual(len(self.rows(self.pending)), 2)
            else:
                self.assertEqual(self.rows(self.invalid)[0]['url'], self.url)
                self.assertEqual(self.rows(self.invalid)[0]['reason'], 'Submission failed')
                self.assertEqual(len(self.rows(self.pending)), 1)

    async def test_invalid_url_leaves_pending_untouched(self):
        for url in (None, '', 'null'):
            with self.assertRaisesRegex(ValueError, 'valid job URL'):
                move_failed_job(self.pending, self.invalid, url, 'Closed')
        self.assertFalse(self.invalid.exists())
        self.assertEqual(len(self.rows(self.pending)), 2)

    async def test_failed_invalid_save_keeps_pending(self):
        with patch('openpyxl.workbook.workbook.Workbook.save', side_effect=OSError('Disk full')):
            with self.assertRaises(OSError):
                move_failed_job(self.pending, self.invalid, self.url, 'Closed')
        self.assertEqual(len(self.rows(self.pending)), 2)

    async def test_retry_after_pending_removal_failure_does_not_duplicate(self):
        with patch('workbook.invalid_jobs.remove_pending_job_url', side_effect=OSError('Locked')):
            with self.assertRaises(OSError):
                move_failed_job(self.pending, self.invalid, self.url, 'Closed')
        move_failed_job(self.pending, self.invalid, self.url, 'Closed')
        self.assertEqual(len(self.rows(self.invalid)), 1)
        self.assertEqual(len(self.rows(self.pending)), 1)
