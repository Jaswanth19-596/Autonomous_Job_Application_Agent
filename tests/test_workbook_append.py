import ast
import tempfile
from types import SimpleNamespace
from unittest.mock import Mock
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from workbook.job_identity import canonical_job_url, canonical_job_id
from workbook.jobs_workbook import upsert_job_row


class WorkbookAppendTests(unittest.TestCase):
    def test_workflow_persists_real_urls_before_generating_identity(self):
        # Isolate the real node from browser/model startup dependencies.
        source = Path(__file__).resolve().parents[1] / "nodes.py"
        node = next(n for n in ast.parse(source.read_text()).body
                    if isinstance(n, ast.FunctionDef) and n.name == "save_job_to_workbook")
        node.decorator_list = []
        writer = Mock(return_value=SimpleNamespace(inserted=1))
        namespace = dict(
            Path=Path, canonical_job_url=canonical_job_url, canonical_job_id=canonical_job_id,
            JOB_COLUMNS={"id", "link", "applyUrl", "companyName"},
            log_node_stage=Mock(), upsert_job_rows=writer,
            pd=SimpleNamespace(Timestamp=SimpleNamespace(now=lambda: SimpleNamespace(strftime=lambda _: "today"))),
        )
        exec(compile(ast.Module(body=[node], type_ignores=[]), str(source), "exec"), namespace)
        state = dict(job_details={"id": "url:74234e98afe7498fb5da", "applyUrl": "null", "link": "null"},
                     job_url="https://example.com/jobs/123", actual_job_url="https://example.com/apply/123",
                     applied_jobs_path="unused.xlsx")
        result = namespace["save_job_to_workbook"](state)["job_details"]
        self.assertEqual(result["link"], state["job_url"])
        self.assertEqual(result["applyUrl"], state["actual_job_url"])
        self.assertEqual(result["id"], canonical_job_id({"applyUrl": state["actual_job_url"]}))
        state["job_url"] = state["actual_job_url"] = "null"
        with self.assertRaisesRegex(ValueError, "stable job ID"):
            namespace["save_job_to_workbook"](state)
        self.assertEqual(writer.call_count, 1)

    def test_placeholder_urls_cannot_identify_a_job(self):
        for value in (None, "", "null", "NULL", "None", "nan", "N/A", "undefined"):
            with self.subTest(value=value):
                self.assertEqual(canonical_job_url(value), "")
                self.assertEqual(canonical_job_id({"applyUrl": value}), "")
                self.assertEqual(
                    canonical_job_id({"applyUrl": value, "link": "https://example.com/jobs/123"}),
                    canonical_job_id({"link": "https://example.com/jobs/123"}),
                )

    def test_new_jobs_append_after_row_55_and_retries_update_in_place(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "jobs.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["id", "applyUrl", "link", "companyName", "reason"])
            for i in range(2, 56):
                sheet.append([f"existing:{i}", "null" if i == 26 else None, None, f"Company {i}", "Keep"])
            book.save(path)
            for i, company in enumerate(("Clasp", "Another company"), 56):
                record = {"link": f"https://example.com/jobs/{i}", "applyUrl": "null", "companyName": company}
                record["id"] = canonical_job_id(record)
                self.assertEqual(upsert_job_row(path, record), ("inserted", i))
                self.assertEqual(upsert_job_row(path, record), ("updated", i))
            result = load_workbook(path)
            self.assertEqual(result.active.max_row, 57)
            self.assertEqual(result.active.cell(26, 4).value, "Company 26")
            self.assertEqual(result.active.cell(26, 5).value, "Keep")
            self.assertIsNone(result.active.cell(56, 5).value)
            result.close()
