import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from jobs.fetcher import JOB_COLUMNS
from models.schemas import JobDetailsModel
import nodes


class JobDetailsWorkbookTests(unittest.TestCase):
    def test_saves_only_current_model_fields_and_required_workflow_metadata(self):
        details = JobDetailsModel(
            title="Backend Engineer",
            companyName="Example Co",
            applyUrl="https://jobs.example.com/roles/123",
        )
        write_result = type("WriteResult", (), {"inserted": 1, "updated": 0})()

        with patch.object(nodes, "upsert_job_rows", return_value=write_result) as upsert:
            nodes.save_job_to_workbook({
                "job_details": details,
                "applied_jobs_path": "/tmp/applied_jobs.xlsx",
                "job_url": "https://jobs.example.com/roles/123",
            })

        row = upsert.call_args.args[1][0]
        model_fields = set(JobDetailsModel.model_fields)
        workflow_metadata = {
            "id",
            "application_status",
            "fetched_at",
            "application_started_at",
            "application_finished_at",
            "application_duration_seconds",
            "application_tool_calls",
            "application_cost_usd",
        }
        self.assertEqual(set(row) - workflow_metadata, {"title", "companyName", "applyUrl"})
        self.assertNotIn("link", row)
        self.assertTrue((set(row) - workflow_metadata).issubset(model_fields))
        self.assertEqual(row["title"], "Backend Engineer")
        self.assertEqual(row["companyName"], "Example Co")

    def test_removed_model_fields_remain_blank_in_the_workbook(self):
        details = JobDetailsModel(
            title="Backend Engineer",
            companyName="Example Co",
            applyUrl="https://jobs.example.com/roles/123",
        )
        with TemporaryDirectory() as directory:
            path = Path(directory) / "applied_jobs.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Applications"
            sheet.append(JOB_COLUMNS)
            workbook.save(path)

            nodes.save_job_to_workbook({
                "job_details": details,
                "applied_jobs_path": str(path),
                "job_url": "https://jobs.example.com/roles/123",
            })

            saved = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet = saved["Applications"]
                headers = [cell.value for cell in sheet[1]]
                values = [cell.value for cell in sheet[2]]
                row = dict(zip(headers, values))
            finally:
                saved.close()

        self.assertEqual(row["title"], "Backend Engineer")
        self.assertEqual(row["companyName"], "Example Co")
        self.assertEqual(row["applyUrl"], "https://jobs.example.com/roles/123")
        self.assertIsNone(row["link"])
        self.assertIsNone(row["salary"])
        self.assertIsNone(row["descriptionText"])

    def test_moves_a_saved_pending_job_into_applied_jobs(self):
        details = JobDetailsModel(
            title="Backend Engineer",
            companyName="Example Co",
            applyUrl="https://jobs.example.com/roles/123",
        )
        pending_url = "https://jobright.ai/jobs/info/example-job"

        with TemporaryDirectory() as directory:
            applied_path = Path(directory) / "applied_jobs.xlsx"
            pending_path = Path(directory) / "pending_jobs.xlsx"

            applied_workbook = Workbook()
            applied_sheet = applied_workbook.active
            applied_sheet.title = "Applications"
            applied_sheet.append(JOB_COLUMNS)
            applied_workbook.save(applied_path)

            pending_workbook = Workbook()
            pending_sheet = pending_workbook.active
            pending_sheet.append(["url", "fetched_at"])
            pending_sheet.append([pending_url, "2026-09-09 10:00:00"])
            pending_sheet.append(["https://jobright.ai/jobs/info/keep", "2026-09-09 10:01:00"])
            pending_workbook.save(pending_path)

            nodes.save_job_to_workbook({
                "job_details": details,
                "applied_jobs_path": str(applied_path),
                "pending_jobs_path": str(pending_path),
                "job_url": pending_url,
            })

            saved_pending_workbook = load_workbook(pending_path, read_only=True, data_only=True)
            try:
                saved_pending_rows = list(
                    saved_pending_workbook.active.iter_rows(min_row=2, values_only=True)
                )
            finally:
                saved_pending_workbook.close()

        self.assertEqual(
            saved_pending_rows,
            [("https://jobright.ai/jobs/info/keep", "2026-09-09 10:01:00")],
        )


if __name__ == "__main__":
    unittest.main()
