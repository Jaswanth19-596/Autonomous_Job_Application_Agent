import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from utilities.application_metrics import (
    empty_metrics,
    record_human_intervention,
    record_llm_usage,
    record_node_timing,
    record_retry,
    record_tool_call,
    stamp_applied_at,
)
from utilities.application_timing import finish_application_timing, start_application_timing


class ApplicationTimingTests(unittest.TestCase):
    def test_duration_and_job_are_persisted_and_timer_resets(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'timings.jsonl'
            with patch('utilities.application_timing.TIMING_LOG_PATH', path), \
                 patch('utilities.application_timing.time.time', side_effect=[100, 112.3456, 200]), \
                 patch('utilities.application_timing.log_workflow_event'):
                state = {**start_application_timing(), 'job_url': 'https://example.com/job',
                         'job_details': {'id': '42', 'title': 'Engineer', 'companyName': 'Example'}}
                result = finish_application_timing(state, 'Applied')
                record = json.loads(path.read_text())
                self.assertEqual(record['elapsed_seconds'], 12.346)
                self.assertEqual(record['job_id'], '42')
                self.assertEqual(record['outcome'], 'Applied')
                self.assertIsNone(result['application_started_at'])
                self.assertEqual(start_application_timing()['application_started_at'], 200)

    def test_rejected_job_does_not_inherit_previous_job_details(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'timings.jsonl'
            with patch('utilities.application_timing.TIMING_LOG_PATH', path), \
                 patch('utilities.application_timing.log_workflow_event'):
                finish_application_timing({'application_started_at': 100,
                    'validation_phase': 'NOT_ELIGIBLE', 'job_url': 'current',
                    'job_details': {'id': 'previous'}}, 'Not Applied')
                record = json.loads(path.read_text())
                self.assertEqual(record['job_url'], 'current')
                self.assertIsNone(record['job_id'])

    def test_legacy_checkpoint_has_no_invented_duration(self):
        self.assertEqual(finish_application_timing({}, 'Applied'), {})

    def test_timing_is_saved_on_matching_excel_row(self):
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "applied_jobs.xlsx"
            book = Workbook()
            book.active.append(["id", "application_status"])
            book.active.append(["42", "Applied"])
            book.active.append(["other", "Not Applied"])
            book.save(workbook_path)
            book.close()
            with patch('utilities.application_timing.TIMING_LOG_PATH', Path(directory) / 'timings.jsonl'), \
                 patch('utilities.application_timing.time.time', return_value=1700000012.3456), \
                 patch('utilities.application_timing.log_workflow_event'):
                finish_application_timing({"application_started_at": 1700000000,
                    "job_details": {"id": "42"}}, "Applied", workbook_path=workbook_path)
            book = load_workbook(workbook_path)
            try:
                rows = list(book.active.values)
                saved = dict(zip(rows[0], rows[1]))
                other = dict(zip(rows[0], rows[2]))
                self.assertEqual(saved["application_elapsed_seconds"], 12.346)
                self.assertEqual(saved["application_status"], "Applied")
                self.assertTrue(saved["application_started_at"].endswith("+00:00"))
                self.assertIsNone(other["application_elapsed_seconds"])
            finally:
                book.close()

    def test_metrics_are_persisted_in_jsonl_and_workbook(self):
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as directory:
            jsonl_path = Path(directory) / "timings.jsonl"
            workbook_path = Path(directory) / "applied_jobs.xlsx"

            book = Workbook()
            book.active.append(["id", "application_status"])
            book.active.append(["99", "Applied"])
            book.save(workbook_path)
            book.close()

            metrics = empty_metrics()
            record_llm_usage(metrics, "muse-spark-1.3-contributor", 1000, 200)
            record_tool_call(metrics, "fill")
            record_tool_call(metrics, "click")
            record_node_timing(metrics, "is_job_valid", 1.5)
            record_human_intervention(metrics)
            stamp_applied_at(metrics)

            state = {
                "application_started_at": 1700000000,
                "job_details": {"id": "99", "title": "Developer", "companyName": "TechCorp"},
                "job_url": "https://example.com/jobs/99",
                "application_metrics": metrics,
            }

            with patch("utilities.application_timing.TIMING_LOG_PATH", jsonl_path), \
                 patch("utilities.application_timing.time.time", return_value=1700000010.5), \
                 patch("utilities.application_timing.log_workflow_event"):
                result = finish_application_timing(state, "Applied", workbook_path=workbook_path)

            # 1. Check returned state reset
            self.assertIsNone(result["application_started_at"])
            self.assertEqual(result["application_elapsed_seconds"], 10.5)
            self.assertEqual(result["application_metrics"], {})

            # 2. Check JSONL content
            record = json.loads(jsonl_path.read_text())
            self.assertEqual(record["total_input_tokens"], 1000)
            self.assertEqual(record["total_output_tokens"], 200)
            self.assertEqual(record["total_tokens"], 1200)
            self.assertEqual(record["llm_call_count"], 1)
            # Cost: 1000 * 0.003 / 1000 + 200 * 0.015 / 1000 = 0.003 + 0.003 = 0.006
            self.assertAlmostEqual(record["total_cost_usd"], 0.006, places=4)
            self.assertEqual(record["tool_call_count"], 2)
            self.assertEqual(record["node_timings"]["is_job_valid"], 1.5)
            self.assertEqual(record["human_intervention_count"], 1)
            self.assertIsNotNone(record["applied_at"])

            # 3. Verify user exclusions: NO tool_calls_by_name and NO human_wait_seconds
            self.assertNotIn("tool_calls_by_name", record)
            self.assertNotIn("human_wait_seconds", record)

            # 4. Check workbook columns populated
            book = load_workbook(workbook_path)
            try:
                rows = list(book.active.values)
                headers = rows[0]
                saved = dict(zip(headers, rows[1]))
                self.assertAlmostEqual(saved["application_cost_usd"], 0.006, places=4)
                self.assertEqual(saved["application_tool_calls"], 2)
                self.assertEqual(saved["llm_call_count"], 1)
                self.assertEqual(saved["total_tokens"], 1200)
                self.assertEqual(saved["human_intervention_count"], 1)
                self.assertTrue(bool(saved["applied_at"]))
            finally:
                book.close()

    def test_application_metrics_accumulators(self):
        m = empty_metrics()
        self.assertEqual(m["total_input_tokens"], 0)
        self.assertEqual(m["total_output_tokens"], 0)
        self.assertEqual(m["total_cost_usd"], 0.0)
        self.assertEqual(m["tool_call_count"], 0)
        self.assertNotIn("tool_calls_by_name", m)
        self.assertNotIn("human_wait_seconds", m)

        # Token + cost accumulation across multiple calls
        record_llm_usage(m, "muse-spark-1.3-contributor", 500, 100)
        record_llm_usage(m, "gpt-5.6-luna", 1000, 500)
        self.assertEqual(m["total_input_tokens"], 1500)
        self.assertEqual(m["total_output_tokens"], 600)
        self.assertEqual(m["total_tokens"], 2100)
        self.assertEqual(m["llm_call_count"], 2)
        # Call 1: 500*0.003/1000 + 100*0.015/1000 = 0.0015 + 0.0015 = 0.003
        # Call 2: 1000*0.005/1000 + 500*0.020/1000 = 0.005 + 0.010 = 0.015
        # Total: 0.018
        self.assertAlmostEqual(m["total_cost_usd"], 0.018, places=4)

        # Tool calls count only
        record_tool_call(m, "take_snapshot")
        record_tool_call(m, "take_snapshot")
        record_tool_call(m, "fill")
        self.assertEqual(m["tool_call_count"], 3)

        # Node timings accumulate
        record_node_timing(m, "extract_job_details", 2.5)
        record_node_timing(m, "extract_job_details", 1.5)
        self.assertEqual(m["node_timings"]["extract_job_details"], 4.0)

        # Retries & human intervention
        record_human_intervention(m)
        record_retry(m)
        self.assertEqual(m["human_intervention_count"], 1)
        self.assertEqual(m["retry_count"], 1)

        # Applied at
        self.assertIsNone(m["applied_at"])
        stamp_applied_at(m)
        self.assertIsNotNone(m["applied_at"])
