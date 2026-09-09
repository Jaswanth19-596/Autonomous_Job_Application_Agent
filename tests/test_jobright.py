import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import AsyncMock, patch

from openpyxl import Workbook, load_workbook

from jobs.jobright import (
    PENDING_JOB_COLUMNS,
    normalize_pending_jobs_workbook,
    persist_jobright_jobs,
    recent_workbook_jobright_urls,
    remove_pending_job_url,
)
from models import (
    ApplicationUrlResolution,
    ApplicationVerificationResult,
    AuthenticationResult,
    JobrightJob,
)
from models.schemas import (
    ApplicationAppliedSubModel,
    ApplicationNotAppliedSubModel,
    ApplicationResult,
)
import nodes
from prompts import load_prompt
class JobrightWorkbookRowsTests(unittest.TestCase):
    def test_persists_only_url_and_fetch_time(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "pending_jobs.xlsx"
            result = persist_jobright_jobs(path, [
                {"link": "https://jobright.ai/jobs/info/first"},
                {"link": "https://jobright.ai/jobs/info/second"},
            ])

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                headers = [cell.value for cell in sheet[1]]
                rows = list(sheet.iter_rows(min_row=2, values_only=True))
            finally:
                workbook.close()

        self.assertEqual(result.inserted, 2)
        self.assertEqual(headers, list(PENDING_JOB_COLUMNS))
        self.assertEqual([row[0] for row in rows], [
            "https://jobright.ai/jobs/info/first",
            "https://jobright.ai/jobs/info/second",
        ])
        self.assertTrue(all(row[1] is not None for row in rows))

    def test_removes_pending_job_by_canonical_url(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "pending_jobs.xlsx"
            persist_jobright_jobs(path, [
                {"link": "https://jobright.ai/jobs/info/remove?utm_source=test"},
                {"link": "https://jobright.ai/jobs/info/keep"},
            ])

            removed = remove_pending_job_url(
                path,
                "https://jobright.ai/jobs/info/remove",
            )
            remaining = recent_workbook_jobright_urls(path)

        self.assertEqual(removed, 1)
        self.assertEqual(remaining, ["https://jobright.ai/jobs/info/keep"])

    def test_normalizes_legacy_pending_rows_to_the_two_column_schema(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "pending_jobs.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["url", "fetched_at", "link", "title"])
            sheet.append([
                "",
                "2026-09-09 10:00:00",
                "https://jobright.ai/jobs/info/legacy",
                "Legacy title",
            ])
            workbook.save(path)

            normalized_count = normalize_pending_jobs_workbook(path)
            saved = load_workbook(path, read_only=True, data_only=True)
            try:
                headers = [cell.value for cell in saved.active[1]]
                rows = list(saved.active.iter_rows(min_row=2, values_only=True))
            finally:
                saved.close()

        self.assertEqual(normalized_count, 1)
        self.assertEqual(headers, list(PENDING_JOB_COLUMNS))
        self.assertEqual(rows, [("https://jobright.ai/jobs/info/legacy", "2026-09-09 10:00:00")])


class JobrightPromptTests(unittest.TestCase):
    def test_single_job_collection_prompts_include_iteration_and_exclusions(self):
        values = {
            "job_number": 1,
            "job_count": 1,
            "jobright_url": "https://jobright.com/jobs/recommended",
            "excluded_jobs": "- https://jobright.com/jobs/info/example-role",
        }

        self.assertIn(
            "exactly one Jobright Top Matched job URL",
            load_prompt("fetch_jobright_jobs/system.txt", **values),
        )
        self.assertIn(
            "job 1 of 1",
            load_prompt("fetch_jobright_jobs/user.txt", **values),
        )
        self.assertIn(
            "selected URL\n(1 of 1)",
            load_prompt("fetch_jobright_jobs/structured_system.txt", **values),
        )
        self.assertIn(
            "URLs must not be fetched again.",
            load_prompt("fetch_jobright_jobs/system.txt", **values),
        )
        self.assertIn(
            "Do not open the detail page, click Apply",
            load_prompt("fetch_jobright_jobs/system.txt", **values),
        )

    def test_returns_recent_jobright_urls_without_using_job_details(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "pending_jobs.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["url", "fetched_at"])
            sheet.append(["https://jobright.com/jobs/info/first", "2026-09-09 10:00:00"])
            sheet.append(["https://jobright.com/jobs/info/second", "2026-09-09 10:01:00"])
            workbook.save(path)

            urls = recent_workbook_jobright_urls(path)

        self.assertEqual(
            urls,
            [
                "https://jobright.com/jobs/info/second",
                "https://jobright.com/jobs/info/first",
            ],
        )


class BrowserAgentTests(unittest.IsolatedAsyncioTestCase):
    def test_verification_prompt_checks_gmail_before_the_application_url(self):
        result = nodes.build_verify_job_status_messages(
            {
                "url_applied_on": "https://jobs.example.com/confirmation/123",
                "is_applied": True,
                "job_details": {"companyName": "Example Co", "title": "Data Engineer"},
            }
        )

        system_message, user_message = result["verify_job_status_messages"]
        self.assertIn("Gmail tools", system_message["content"])
        self.assertIn("Example Co", user_message["content"])
        self.assertIn("Data Engineer", user_message["content"])
        self.assertLess(
            system_message["content"].index("First use the Gmail tools"),
            system_message["content"].index("Only when Gmail has no qualifying"),
        )

    async def test_verification_accepts_matching_confirmation_email(self):
        verifier_model = AsyncMock()
        verifier_model.ainvoke.return_value = type("Response", (), {"tool_calls": []})()
        structured_model = AsyncMock()
        structured_model.ainvoke.return_value = ApplicationVerificationResult(
            application_status="Applied",
            evidence_source="confirmation_email",
            message="Example Co confirmed receipt of the Data Engineer application.",
        )

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        with (
            patch.object(
                nodes,
                "get_custom_model_with_tools",
                new_callable=AsyncMock,
                return_value=verifier_model,
            ) as get_model,
            patch.object(nodes, "get_custom_structured_model", return_value=structured_model),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            result = await nodes.verify_job_status(
                {"verify_job_status_messages": [{"role": "user", "content": "Verify."}]}
            )

        get_model.assert_awaited_once_with(["chrome", "gmail"])
        self.assertEqual(result["verification_phase"], "APPLIED")
        self.assertEqual(result["application_status"], "Applied")
        self.assertTrue(result["is_applied"])
        self.assertIn("confirmed receipt", result["final_status_message"])

    async def test_dispatch_resets_authentication_state_for_the_next_job(self):
        result = await nodes.dispath_job(
            {
                "application_queue": [
                    {"id": "job-2", "applyUrl": "https://jobs.example.com/apply/2"}
                ],
                "authenticated_session": True,
                "actual_job_url": "https://jobs.example.com/apply/1",
                "resolve_application_url_messages": [
                    {"role": "user", "content": "Old login evidence."}
                ],
                "perform_login_messages": [
                    {"role": "user", "content": "Old login evidence."}
                ],
            }
        )

        self.assertEqual(result["job_url"], "https://jobs.example.com/apply/2")
        self.assertFalse(result["authenticated_session"])
        self.assertEqual(result["actual_job_url"], "")
        self.assertEqual(result["resolve_application_url_messages"], [])
        self.assertEqual(result["perform_login_messages"], [])

    async def test_submission_node_records_the_successful_application_branch(self):
        browser_model = AsyncMock()
        browser_model.ainvoke.return_value = type("Response", (), {"tool_calls": []})()
        structured_model = AsyncMock()
        structured_model.ainvoke.return_value = ApplicationResult(
            decision=ApplicationAppliedSubModel(
                final_status_message="Application submitted successfully.",
                url_applied_on="https://jobs.example.com/confirmation/123",
            )
        )

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        with (
            patch.object(
                nodes,
                "get_custom_model_with_tools",
                new_callable=AsyncMock,
                return_value=browser_model,
            ),
            patch.object(nodes, "get_custom_structured_model", return_value=structured_model),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            result = await nodes.submit_job_application(
                {"submit_job_application_messages": [{"role": "user", "content": "Apply."}]}
            )

        self.assertEqual(result["job_application_phase"], "APPLIED")
        self.assertTrue(result["is_applied"])
        self.assertEqual(result["url_applied_on"], "https://jobs.example.com/confirmation/123")
        self.assertFalse(result["requires_human"])

    async def test_submission_node_emits_only_the_new_tool_request(self):
        tool_call = {"name": "take_snapshot", "args": {}, "id": "call-1"}
        browser_model = AsyncMock()
        browser_model.ainvoke.return_value = type("Response", (), {"tool_calls": [tool_call]})()

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        prior_messages = [
            {"role": "system", "content": "Use fill_form to replace form values."},
            {"role": "user", "content": "Apply to this job."},
        ]
        with (
            patch.object(nodes, "get_custom_model_with_tools", new_callable=AsyncMock, return_value=browser_model),
            patch.object(nodes, "get_custom_structured_model"),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            result = await nodes.submit_job_application(
                {"submit_job_application_messages": prior_messages}
            )

        self.assertEqual(result["job_application_phase"], "NEEDS_TOOLS")
        self.assertEqual(len(result["submit_job_application_messages"]), 1)
        self.assertIs(result["submit_job_application_messages"][0], browser_model.ainvoke.return_value)

    async def test_application_tool_node_blocks_email_keyboard_typing(self):
        tool_call = {
            "name": "type_text",
            "args": {"text": "madhajaswanth91@gmail.com"},
            "id": "call-email",
        }
        response = type("Response", (), {"tool_calls": [tool_call]})()

        with (
            patch.object(nodes, "log_tool_calls"),
            patch.object(nodes, "log_tool_result"),
        ):
            result = await nodes.tool_node(
                {
                    "tool_call_origin": "submit_job_application",
                    "submit_job_application_messages": [response],
                }
            )

        tool_message = result["submit_job_application_messages"][0]
        self.assertIn("Blocked unsafe email keyboard entry", tool_message.content)
        self.assertEqual(tool_message.tool_call_id, "call-email")

    async def test_submission_node_records_the_unsuccessful_human_branch(self):
        browser_model = AsyncMock()
        browser_model.ainvoke.return_value = type("Response", (), {"tool_calls": []})()
        structured_model = AsyncMock()
        structured_model.ainvoke.return_value = ApplicationResult(
            decision=ApplicationNotAppliedSubModel(
                requires_human=True,
                human_intervention_reason="CAPTCHA",
                human_intervention_prompt="Complete the CAPTCHA on the application page.",
            )
        )

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        with (
            patch.object(
                nodes,
                "get_custom_model_with_tools",
                new_callable=AsyncMock,
                return_value=browser_model,
            ),
            patch.object(nodes, "get_custom_structured_model", return_value=structured_model),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            result = await nodes.submit_job_application(
                {"submit_job_application_messages": [{"role": "user", "content": "Apply."}]}
            )

        self.assertEqual(result["job_application_phase"], "NOT_APPLIED")
        self.assertFalse(result["is_applied"])
        self.assertTrue(result["requires_human"])
        self.assertEqual(result["human_intervention_reason"], "CAPTCHA")
        self.assertEqual(
            result["human_intervention_prompt"],
            "Complete the CAPTCHA on the application page.",
        )

    def test_application_result_is_discriminated_by_is_applied(self):
        applied = ApplicationResult.model_validate(
            {
                "decision": {
                    "is_applied": True,
                    "final_status_message": "Application submitted successfully.",
                    "url_applied_on": "https://jobs.example.com/confirmation/123",
                }
            }
        )
        not_applied = ApplicationResult.model_validate(
            {
                "decision": {
                    "is_applied": False,
                    "requires_human": False,
                    "human_intervention_reason": None,
                    "human_intervention_prompt": "",
                }
            }
        )

        self.assertIsInstance(applied.decision, ApplicationAppliedSubModel)
        self.assertIsInstance(not_applied.decision, ApplicationNotAppliedSubModel)

    async def test_resolution_node_hands_the_ready_url_to_the_submission_node(self):
        browser_model = AsyncMock()
        browser_model.ainvoke.return_value = type("Response", (), {"tool_calls": []})()
        structured_model = AsyncMock()
        structured_model.ainvoke.return_value = ApplicationUrlResolution(
            ready_to_apply=True,
            actual_job_url="https://jobs.example.com/apply/123",
            message="Application form is ready after sign-in.",
        )

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        with (
            patch.object(
                nodes,
                "get_custom_model_with_tools",
                new_callable=AsyncMock,
                return_value=browser_model,
            ) as get_model,
            patch.object(nodes, "get_custom_structured_model", return_value=structured_model),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            result = await nodes.resolve_application_url(
                {"job_url": "https://jobs.example.com/posting/123"}
            )

        get_model.assert_awaited_once_with(["chrome"])
        self.assertEqual(result["resolution_phase"], "SUCCESS")
        self.assertEqual(result["actual_job_url"], "https://jobs.example.com/apply/123")
        self.assertEqual(
            result["submit_job_application_messages"][0]["content"],
            "The Actual Job URL is : https://jobs.example.com/apply/123\n",
        )

    async def test_resolution_node_routes_browser_calls_to_the_shared_tool_node(self):
        tool_call = {"name": "list_pages", "args": {}, "id": "call-1"}
        browser_model = AsyncMock()
        browser_model.ainvoke.return_value = type("Response", (), {"tool_calls": [tool_call]})()

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        with (
            patch.object(
                nodes,
                "get_custom_model_with_tools",
                new_callable=AsyncMock,
                return_value=browser_model,
            ),
            patch.object(nodes, "get_custom_structured_model"),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            result = await nodes.resolve_application_url(
                {"job_url": "https://jobs.example.com/posting/123"}
            )

        self.assertEqual(result["resolution_phase"], "NEEDS_TOOLS")
        self.assertEqual(result["tool_call_origin"], "resolve_application_url")
        self.assertEqual(len(result["resolve_application_url_messages"]), 3)
        self.assertEqual(
            result["resolve_application_url_messages"][0]["role"], "system"
        )
        self.assertEqual(
            result["resolve_application_url_messages"][-1].tool_calls,
            [tool_call],
        )

    async def test_resolution_node_records_a_blocked_login_without_submitting(self):
        browser_model = AsyncMock()
        browser_model.ainvoke.return_value = type("Response", (), {"tool_calls": []})()
        structured_model = AsyncMock()
        structured_model.ainvoke.return_value = ApplicationUrlResolution(
            ready_to_apply=False,
            actual_job_url="https://jobs.example.com/login",
            message="Sign-in was blocked by a CAPTCHA.",
        )

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        with (
            patch.object(
                nodes,
                "get_custom_model_with_tools",
                new_callable=AsyncMock,
                return_value=browser_model,
            ),
            patch.object(nodes, "get_custom_structured_model", return_value=structured_model),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            result = await nodes.resolve_application_url(
                {"job_url": "https://jobs.example.com/posting/123"}
            )

        self.assertEqual(result["resolution_phase"], "FAILED")
        self.assertEqual(result["submit_job_application_messages"], [])

    async def test_resolution_node_delegates_a_login_gate_and_remembers_where_to_return(self):
        browser_model = AsyncMock()
        browser_model.ainvoke.return_value = type("Response", (), {"tool_calls": []})()
        structured_model = AsyncMock()
        structured_model.ainvoke.return_value = ApplicationUrlResolution(
            ready_to_apply=False,
            requires_login=True,
            actual_job_url="https://jobs.example.com/login",
            message="Sign-in is required.",
        )

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        with (
            patch.object(nodes, "get_custom_model_with_tools", new_callable=AsyncMock, return_value=browser_model),
            patch.object(nodes, "get_custom_structured_model", return_value=structured_model),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            result = await nodes.resolve_application_url(
                {"job_url": "https://jobs.example.com/posting/123"}
            )

        self.assertEqual(result["resolution_phase"], "REQUIRES_LOGIN")
        self.assertEqual(result["login_return_node"], "resolve_application_url")
        self.assertEqual(result["actual_job_url"], "https://jobs.example.com/login")

    async def test_resolution_after_successful_login_does_not_retry_login(self):
        browser_model = AsyncMock()
        browser_model.ainvoke.return_value = type("Response", (), {"tool_calls": []})()
        structured_model = AsyncMock()
        structured_model.ainvoke.return_value = ApplicationUrlResolution(
            ready_to_apply=False,
            requires_login=True,
            actual_job_url="https://jobs.example.com/login",
            message="Sign-in is required.",
        )

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        with (
            patch.object(
                nodes,
                "get_custom_model_with_tools",
                new_callable=AsyncMock,
                return_value=browser_model,
            ),
            patch.object(nodes, "get_custom_structured_model", return_value=structured_model),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            result = await nodes.resolve_application_url(
                {
                    "job_url": "https://jobs.example.com/posting/123",
                    "actual_job_url": "https://jobs.example.com/apply/123",
                    "authenticated_session": True,
                    "resolve_application_url_messages": [],
                }
            )

        self.assertEqual(result["resolution_phase"], "FAILED")
        self.assertTrue(result["requires_human"])
        prompt_messages = browser_model.ainvoke.await_args.args[0]
        self.assertIn("https://jobs.example.com/apply/123", prompt_messages[1]["content"])
        self.assertIn("just verified an authenticated session", prompt_messages[2]["content"])

    async def test_authenticated_resolution_preserves_tool_evidence_after_the_first_turn(self):
        tool_call = {"name": "list_pages", "args": {}, "id": "call-1"}
        browser_model = AsyncMock()
        browser_model.ainvoke.side_effect = [
            type("Response", (), {"tool_calls": [tool_call]})(),
            type("Response", (), {"tool_calls": []})(),
        ]
        structured_model = AsyncMock()
        structured_model.ainvoke.return_value = ApplicationUrlResolution(
            ready_to_apply=True,
            requires_login=False,
            actual_job_url="https://jobs.example.com/apply/123",
            message="Application form is ready.",
        )

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        state = {
            "job_url": "https://jobs.example.com/posting/123",
            "actual_job_url": "https://jobs.example.com/apply/123",
            "authenticated_session": True,
            "resolve_application_url_messages": [],
        }
        with (
            patch.object(
                nodes,
                "get_custom_model_with_tools",
                new_callable=AsyncMock,
                return_value=browser_model,
            ),
            patch.object(nodes, "get_custom_structured_model", return_value=structured_model),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            first_turn = await nodes.resolve_application_url(state)
            second_turn = await nodes.resolve_application_url(
                {
                    **state,
                    "resolve_application_url_messages": first_turn[
                        "resolve_application_url_messages"
                    ],
                }
            )

        self.assertEqual(first_turn["resolution_phase"], "NEEDS_TOOLS")
        self.assertEqual(second_turn["resolution_phase"], "SUCCESS")
        self.assertEqual(
            browser_model.ainvoke.await_args_list[1].args[0],
            first_turn["resolve_application_url_messages"],
        )

    async def test_submission_node_delegates_an_unexpected_login_and_returns_to_itself(self):
        browser_model = AsyncMock()
        browser_model.ainvoke.return_value = type("Response", (), {"tool_calls": []})()
        structured_model = AsyncMock()
        structured_model.ainvoke.return_value = ApplicationResult(
            decision=ApplicationNotAppliedSubModel(
                requires_login=True,
                requires_human=False,
                human_intervention_reason=None,
                human_intervention_prompt="",
            )
        )

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        with (
            patch.object(nodes, "get_custom_model_with_tools", new_callable=AsyncMock, return_value=browser_model),
            patch.object(nodes, "get_custom_structured_model", return_value=structured_model),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            result = await nodes.submit_job_application(
                {"submit_job_application_messages": [{"role": "user", "content": "Apply."}]}
            )

        self.assertEqual(result["job_application_phase"], "REQUIRES_LOGIN")
        self.assertEqual(result["login_return_node"], "submit_job_application")

    async def test_login_node_returns_a_successful_login_to_its_caller(self):
        browser_model = AsyncMock()
        browser_model.ainvoke.return_value = type("Response", (), {"tool_calls": []})()
        structured_model = AsyncMock()
        structured_model.ainvoke.return_value = AuthenticationResult(
            is_logged_in=True,
            logged_in_url="https://jobs.example.com/apply/123",
            message="Signed in.",
        )

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        with (
            patch.object(nodes, "get_custom_model_with_tools", new_callable=AsyncMock, return_value=browser_model) as get_model,
            patch.object(nodes, "get_custom_structured_model", return_value=structured_model),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            result = await nodes.perform_login(
                {
                    "actual_job_url": "https://jobs.example.com/login",
                    "login_return_node": "submit_job_application",
                }
            )

        get_model.assert_awaited_once_with(["chrome", "gmail"])
        self.assertEqual(result["login_phase"], "SUCCESS")
        self.assertEqual(result["actual_job_url"], "https://jobs.example.com/apply/123")
        self.assertTrue(result["authenticated_session"])
        self.assertEqual(result["resolve_application_url_messages"], [])
        self.assertEqual(result["perform_login_messages"], [])

    async def test_fetch_node_routes_browser_calls_to_the_shared_tool_node(self):
        tool_call = {"name": "navigate_page", "args": {"url": "https://jobright.com"}, "id": "call-1"}
        model = AsyncMock()
        model.ainvoke.return_value = type("Response", (), {"tool_calls": [tool_call]})()

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        with (
            patch.object(nodes, "recent_workbook_jobright_urls", return_value=[]) as recent_workbook_urls,
            patch.object(
                nodes,
                "get_custom_model_with_tools",
                new_callable=AsyncMock,
                return_value=model,
            ),
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
            patch.object(nodes, "log_tool_calls"),
        ):
            result = await nodes.fetch_jobright_jobs(
                {"jobright_job_count": 1, "pending_jobs_path": "/tmp/pending_jobs.xlsx"}
            )

        recent_workbook_urls.assert_called_once_with(Path("/tmp/pending_jobs.xlsx"), limit=20)
        self.assertEqual(result["next_step"], "tool_node")
        self.assertEqual(result["tool_call_origin"], "fetch_jobright_jobs")
        self.assertEqual(result["jobright_collection_messages"][-1].tool_calls, [tool_call])

    async def test_shared_tool_node_appends_jobright_tool_results_to_its_transcript(self):
        tool_call = {"name": "navigate_page", "args": {"url": "https://jobright.com"}, "id": "call-1"}
        browser_tool = AsyncMock()
        browser_tool.ainvoke.return_value = "Jobright recommendations are visible."
        response = type("Response", (), {"tool_calls": [tool_call]})()

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable

        original_tools = nodes.tools_by_name.copy()
        nodes.tools_by_name["navigate_page"] = browser_tool
        try:
            with (
                patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
                patch.object(nodes, "log_tool_calls"),
                patch.object(nodes, "log_tool_result"),
            ):
                result = await nodes.tool_node({
                    "tool_call_origin": "fetch_jobright_jobs",
                    "jobright_collection_messages": [
                        {"role": "system", "content": "browser instructions"},
                        response,
                    ],
                })
        finally:
            nodes.tools_by_name.clear()
            nodes.tools_by_name.update(original_tools)

        browser_tool.ainvoke.assert_awaited_once_with({"url": "https://jobright.com"})
        self.assertEqual(len(result["jobright_collection_messages"]), 3)
        self.assertIsInstance(result["jobright_collection_messages"][-1], nodes.ToolMessage)

    async def test_fetch_node_persists_a_completed_job_through_jobright_helper(self):
        browser_model = AsyncMock()
        browser_model.ainvoke.return_value = type("Response", (), {"tool_calls": []})()
        structured_model = AsyncMock()
        structured_model.ainvoke.return_value = JobrightJob(
            link="https://jobright.com/jobs/info/job-1",
        )

        async def pass_through_progress(_node, _activity, awaitable, **_kwargs):
            return await awaitable
        with (
            patch.object(
                nodes,
                "get_custom_model_with_tools",
                new_callable=AsyncMock,
                return_value=browser_model,
            ),
            patch.object(nodes, "get_custom_structured_model", return_value=structured_model),
            patch.object(nodes, "persist_jobright_jobs") as persist,
            patch.object(nodes, "await_with_progress", side_effect=pass_through_progress),
        ):
            persist.return_value = type("WriteResult", (), {"inserted": 1, "updated": 0})()
            result = await nodes.fetch_jobright_jobs(
                {
                    "jobright_job_count": 1,
                    "pending_jobs_path": "/tmp/pending_jobs.xlsx",
                    "jobright_excluded_jobs": [],
                    "jobright_collection_messages": [
                        {"role": "system", "content": "browser instructions"},
                        {"role": "user", "content": "browser task"},
                        nodes.ToolMessage(content="snapshot", tool_call_id="call-1", name="snapshot"),
                    ],
                }
            )

        persist.assert_called_once()
        self.assertEqual(len(result["jobright_jobs"]), 1)
        self.assertEqual(result["next_step"], "collect_user_request")


if __name__ == "__main__":
    unittest.main()
