"""Exercise real workbook operations against an in-memory Drive transport."""

import hashlib
from requests.exceptions import ReadTimeout
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from jobs.fetcher import save_jobs_to_excel
from workbook.drive import DriveClient, DriveStorageError, XLSX_MIME, get_credentials
from workbook.drive_setup import migrate
from workbook.invalid_jobs import move_failed_job
from workbook.jobs_workbook import get_job_row, update_job_row, upsert_job_rows
from workbook.pending_jobs import remove_pending_job_url
from workbook.storage import WORKBOOK_NAMES, workbook_file


def xlsx(rows, *, dashboard=False):
    book = Workbook()
    book.active.title = "Applications"
    for row in rows:
        book.active.append(row)
    book.active["A1"].font = Font(bold=True)
    if dashboard:
        book.create_sheet("Dashboard")["A1"] = "=COUNTA(Applications!A:A)-1"
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


class FakeDrive:
    def __init__(self, files):
        self.files = dict(files)
        self.versions = {file_id: 1 for file_id in files}
        self.uploads = []
        self.downloads = []
        self.fail_upload = False
        self.fail_download = False

    def download(self, file_id, destination):
        if self.fail_download:
            raise DriveStorageError("Drive unavailable")
        self.downloads.append(file_id)
        destination.write_bytes(self.files[file_id])
        return str(self.versions[file_id]), f'"{self.versions[file_id]}"'

    def update(self, file_id, source, *, version, etag):
        if self.fail_upload:
            raise DriveStorageError("Upload failed")
        if str(self.versions[file_id]) != version:
            raise DriveStorageError("Drive workbook changed")
        self.files[file_id] = source.read_bytes()
        self.versions[file_id] += 1
        self.uploads.append(file_id)

    def metadata(self, file_id):
        return {"version": str(self.versions[file_id]), "mimeType": XLSX_MIME}, None

    def close(self):
        pass


class DriveStorageTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        (self.root / "data").mkdir()
        self.paths = {name: self.root / "data" / name for name in WORKBOOK_NAMES}
        self.manifest = self.root / "drive.json"
        self.ids = {name: f"file-{index}" for index, name in enumerate(WORKBOOK_NAMES)}
        self.manifest.write_text(json.dumps({"folder_id": "folder", "files": self.ids}))
        self.url = "https://example.com/jobs/42"
        self.drive = FakeDrive({
            self.ids["pending_jobs.xlsx"]: xlsx([["url", "fetched_at"], [self.url, "today"]]),
            self.ids["applied_jobs.xlsx"]: xlsx([["id", "link", "application_status"], ["42", self.url, "Not Applied"]], dashboard=True),
            self.ids["invalid_jobs.xlsx"]: xlsx([["job_id", "url", "reason"]]),
        })
        for name, path in self.paths.items():
            path.write_bytes(xlsx([["stale local data"]]))
        for patcher in (
            patch.dict(os.environ, {"JOB_AGENT_DRIVE_CONFIG": str(self.manifest), "JOB_AGENT_WORKBOOK_STORAGE": "drive"}),
            patch("workbook.storage.PROJECT_ROOT", self.root),
            patch("workbook.storage.get_credentials", return_value=Mock()),
            patch("workbook.storage.DriveClient", return_value=self.drive),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)

    def rows(self, name):
        book = load_workbook(BytesIO(self.drive.files[self.ids[name]]), data_only=True)
        try:
            return list(book.active.values)
        finally:
            book.close()

    def test_read_uses_current_drive_data_and_leaves_local_original_untouched(self):
        path = self.paths["applied_jobs.xlsx"]
        original = path.read_bytes()
        self.assertEqual(get_job_row(path, "42")["application_status"], "Not Applied")
        update_job_row(path, "42", {"application_status": "Applied"})
        self.assertEqual(get_job_row(path, "42")["application_status"], "Applied")
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual(len(self.drive.downloads), 3)

    def test_updates_preserve_dashboard_formula_and_header_format(self):
        update_job_row(self.paths["applied_jobs.xlsx"], "42", {"application_status": "Applied"})
        book = load_workbook(BytesIO(self.drive.files[self.ids["applied_jobs.xlsx"]]))
        try:
            self.assertEqual(book["Dashboard"]["A1"].value, "=COUNTA(Applications!A:A)-1")
            self.assertTrue(book["Applications"]["A1"].font.bold)
        finally:
            book.close()

    def test_batch_downloads_and_uploads_once(self):
        upsert_job_rows(self.paths["applied_jobs.xlsx"], [
            {"id": "43", "link": "https://example.com/jobs/43"},
            {"id": "44", "link": "https://example.com/jobs/44"},
        ])
        self.assertEqual(self.drive.downloads, [self.ids["applied_jobs.xlsx"]])
        self.assertEqual(self.drive.uploads, [self.ids["applied_jobs.xlsx"]])
        self.assertEqual(len(self.rows("applied_jobs.xlsx")), 4)

    def test_fetcher_appends_to_remote_pending_queue(self):
        save_jobs_to_excel(["https://example.com/jobs/43"], self.paths["pending_jobs.xlsx"])
        self.assertEqual([r[0] for r in self.rows("pending_jobs.xlsx")[1:]], [self.url, "https://example.com/jobs/43"])

    def test_noop_queue_removal_does_not_upload(self):
        self.assertEqual(remove_pending_job_url(self.paths["pending_jobs.xlsx"], "https://example.com/jobs/missing"), 0)
        self.assertEqual(self.drive.uploads, [])

    def test_rejected_record_uploads_before_pending_removal(self):
        move_failed_job(self.paths["pending_jobs.xlsx"], self.paths["invalid_jobs.xlsx"], self.url, "Closed")
        self.assertEqual(self.drive.uploads, [self.ids["invalid_jobs.xlsx"], self.ids["pending_jobs.xlsx"]])
        self.assertEqual(len(self.rows("pending_jobs.xlsx")), 1)
        self.assertEqual(self.rows("invalid_jobs.xlsx")[1][2], "Closed")

    def test_failed_rejection_upload_keeps_remote_pending_queue(self):
        before = self.drive.files[self.ids["pending_jobs.xlsx"]]
        self.drive.fail_upload = True
        with self.assertRaisesRegex(DriveStorageError, "Upload failed"):
            move_failed_job(self.paths["pending_jobs.xlsx"], self.paths["invalid_jobs.xlsx"], self.url, "Closed")
        self.assertEqual(self.drive.files[self.ids["pending_jobs.xlsx"]], before)

    def test_failed_upload_preserves_private_recovery_copy(self):
        self.drive.fail_upload = True
        with self.assertRaisesRegex(DriveStorageError, "Unsynced workbook preserved"):
            update_job_row(self.paths["applied_jobs.xlsx"], "42", {"application_status": "Applied"})
        copies = list((self.root / "data/.drive-recovery").glob("*.xlsx"))
        self.assertEqual(len(copies), 1)
        self.assertEqual(copies[0].stat().st_mode & 0o777, 0o600)
        self.assertEqual(get_job_row(copies[0], "42")["application_status"], "Applied")

    def test_drive_failure_never_falls_back_to_local_file(self):
        self.drive.fail_download = True
        with self.assertRaisesRegex(DriveStorageError, "unavailable"):
            get_job_row(self.paths["applied_jobs.xlsx"], "42")

    def test_corrupt_remote_pending_file_is_not_overwritten(self):
        file_id = self.ids["pending_jobs.xlsx"]
        self.drive.files[file_id] = b"not an Excel workbook"
        with self.assertRaisesRegex(ValueError, "refusing to overwrite"):
            save_jobs_to_excel(["https://example.com/jobs/new"], self.paths["pending_jobs.xlsx"])
        self.assertEqual(self.drive.files[file_id], b"not an Excel workbook")
        self.assertEqual(self.drive.uploads, [])

    def test_exception_does_not_upload_and_temporary_file_is_removed(self):
        working = None
        with self.assertRaisesRegex(ValueError, "abort"):
            with workbook_file(self.paths["pending_jobs.xlsx"], write=True) as local:
                working = local
                local.write_bytes(b"incomplete")
                raise ValueError("abort")
        self.assertEqual(self.drive.uploads, [])
        self.assertFalse(working.exists())

    def test_remote_change_is_not_overwritten(self):
        file_id = self.ids["pending_jobs.xlsx"]
        before = self.drive.files[file_id]
        with self.assertRaisesRegex(DriveStorageError, "changed"):
            with workbook_file(self.paths["pending_jobs.xlsx"], write=True) as local:
                local.write_bytes(xlsx([["url"], ["https://example.com/jobs/new"]]))
                self.drive.versions[file_id] += 1
        self.assertEqual(self.drive.files[file_id], before)

    def test_missing_configuration_does_not_use_local_when_drive_required(self):
        self.manifest.unlink()
        with self.assertRaisesRegex(DriveStorageError, "not been configured"):
            get_job_row(self.paths["applied_jobs.xlsx"], "42")

    def test_custom_path_stays_local(self):
        custom = self.root / "custom.xlsx"
        custom.write_bytes(xlsx([["id", "link"], ["custom", self.url]]))
        self.assertEqual(get_job_row(custom, "custom")["id"], "custom")
        self.assertEqual(self.drive.downloads, [])

    def test_tool_reads_remote_data_without_local_files(self):
        from utilities.job_context import get_pending_or_invalid_jobs
        from utilities.job_search import get_applied_jobs
        for path in self.paths.values():
            path.unlink()
        with patch.dict("utilities.job_context._WORKBOOKS", {
            "pending": self.paths["pending_jobs.xlsx"], "invalid": self.paths["invalid_jobs.xlsx"],
        }), patch("utilities.job_search._EXCEL_PATH", self.paths["applied_jobs.xlsx"]):
            pending = get_pending_or_invalid_jobs.invoke({"job_type": "pending"})
            applied = get_applied_jobs.invoke({"job_id": "42"})
        self.assertEqual(pending["jobs"][0]["url"], self.url)
        self.assertEqual(applied["total_matching"], 1)


class DriveTransportTests(unittest.TestCase):
    def setUp(self):
        sleeper = patch("workbook.drive.time.sleep")
        sleeper.start()
        self.addCleanup(sleeper.stop)

    def test_read_timeout_retries_and_exhausts(self):
        client = DriveClient(Mock())
        self.addCleanup(client.close)
        success = Mock(ok=True, status_code=200)
        with patch.object(client.session, "request", side_effect=[ReadTimeout(), success]) as request:
            self.assertIs(client.request("GET", "url"), success)
            self.assertEqual(request.call_count, 2)
        with patch.object(client.session, "request", side_effect=ReadTimeout()) as request:
            with self.assertRaises(DriveStorageError):
                client.request("GET", "url")
            self.assertEqual(request.call_count, 4)

    def test_upload_reconciles_timeout_and_replays_complete_payload(self):
        for committed, changed in ((False, False), (True, False), (False, True)):
            with self.subTest(committed=committed, changed=changed), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "book.xlsx"
                source.write_bytes(b"complete workbook")
                client = DriveClient(Mock())
                self.addCleanup(client.close)
                current = {"version": "2" if committed or changed else "1"}
                if committed:
                    current["md5Checksum"] = hashlib.md5(source.read_bytes()).hexdigest()
                with patch.object(client, "metadata", side_effect=[({"version": "1"}, None), (current, None)]), patch.object(
                    client.session, "request", side_effect=[ReadTimeout(), Mock(ok=True, status_code=200)]
                ) as request:
                    if changed:
                        with self.assertRaisesRegex(DriveStorageError, "changed"):
                            client.update("file", source, version="1", etag='"1"')
                    else:
                        client.update("file", source, version="1", etag='"1"')
                    self.assertEqual(request.call_count, 1 if committed or changed else 2)
                    for call in request.call_args_list:
                        self.assertEqual(call.kwargs["data"], source.read_bytes())
                        self.assertEqual(call.kwargs["headers"]["If-Match"], '"1"')

    def test_create_is_not_blindly_retried(self):
        client = DriveClient(Mock())
        self.addCleanup(client.close)
        with patch.object(client.session, "request", side_effect=ReadTimeout()) as request:
            with self.assertRaises(DriveStorageError):
                client.create_folder("jobs")
            self.assertEqual(request.call_count, 1)

    def test_download_retries_when_google_finalizes_uploaded_file(self):
        client = DriveClient(Mock())
        self.addCleanup(client.close)
        versions = ["1", "2", "2", "2"]
        metadata = [({"version": version, "mimeType": XLSX_MIME}, None) for version in versions]
        with tempfile.TemporaryDirectory() as directory, \
             patch.object(client, "metadata", side_effect=metadata), \
             patch.object(client, "request", side_effect=[Mock(content=b"old"), Mock(content=b"current")]):
            destination = Path(directory) / "workbook.xlsx"
            self.assertEqual(client.download("file", destination), ("2", None))
            self.assertEqual(destination.read_bytes(), b"current")

    def test_update_checks_version_before_sending_content(self):
        client = DriveClient(Mock())
        self.addCleanup(client.close)
        with patch.object(client, "metadata", return_value=({"version": "new"}, None)), \
             patch.object(client, "request") as request:
            with self.assertRaisesRegex(DriveStorageError, "changed"):
                client.update("id", Path("unused.xlsx"), version="old", etag=None)
            request.assert_not_called()

    def test_http_failure_and_conflict_are_errors(self):
        client = DriveClient(Mock())
        self.addCleanup(client.close)
        for status in (403, 404, 412, 500):
            with patch.object(client.session, "request", return_value=Mock(ok=False, status_code=status)):
                with self.assertRaises(DriveStorageError):
                    client.request("GET", "https://www.googleapis.com/drive/v3/files/id")

    def test_runtime_does_not_open_browser_for_missing_authorization(self):
        with tempfile.TemporaryDirectory() as directory, \
             patch.dict(os.environ, {"JOB_AGENT_DRIVE_TOKEN_PATH": str(Path(directory) / "missing.json")}), \
             patch("workbook.drive.InstalledAppFlow") as flow:
            with self.assertRaisesRegex(DriveStorageError, "authorization"):
                get_credentials()
            flow.assert_not_called()


class DriveMigrationTests(unittest.TestCase):
    def test_migration_preserves_originals_and_repeat_setup_uses_drive(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "data").mkdir()
            originals = {}
            for name in WORKBOOK_NAMES:
                originals[name] = xlsx([["id"], [name]])
                (root / "data" / name).write_bytes(originals[name])
            manifest = root / "drive.json"
            client = FakeDrive({})
            client.find = Mock(return_value=[])
            client.create_folder = Mock(return_value={"id": "folder"})
            def create(source, folder_id):
                file_id = source.name
                client.files[file_id] = source.read_bytes()
                client.versions[file_id] = 1
                return {"id": file_id}
            client.create_workbook = Mock(side_effect=create)
            with patch("workbook.drive_setup.PROJECT_ROOT", root), \
                 patch("workbook.drive_setup.config_path", return_value=manifest):
                config = migrate(client)
                self.assertEqual(len(config["files"]), 3)
                self.assertTrue(manifest.is_file())
                self.assertEqual(len(list((root / "data/drive-migration-backups").glob("*/*.xlsx"))), 3)
                (root / "data/applied_jobs.xlsx").write_bytes(b"stale local data")
                self.assertEqual(migrate(client), config)
            self.assertEqual(client.create_workbook.call_count, 3)
            self.assertEqual(client.files["applied_jobs.xlsx"], originals["applied_jobs.xlsx"])

    def test_failed_upload_does_not_activate_drive(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "data").mkdir()
            for name in WORKBOOK_NAMES:
                (root / "data" / name).write_bytes(xlsx([["id"]]))
            manifest = root / "drive.json"
            client = Mock()
            client.find.return_value = []
            client.create_folder.return_value = {"id": "folder"}
            client.create_workbook.side_effect = DriveStorageError("Upload failed")
            with patch("workbook.drive_setup.PROJECT_ROOT", root), \
                 patch("workbook.drive_setup.config_path", return_value=manifest):
                with self.assertRaises(DriveStorageError):
                    migrate(client)
            self.assertFalse(manifest.exists())
            self.assertTrue(all((root / "data" / name).is_file() for name in WORKBOOK_NAMES))

    def test_existing_different_drive_file_is_never_overwritten(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "data").mkdir()
            for name in WORKBOOK_NAMES:
                (root / "data" / name).write_bytes(xlsx([["local"]]))
            manifest = root / "drive.json"
            client = FakeDrive({"remote": xlsx([["edited on Drive"]])})
            client.find = Mock(side_effect=[[{"id": "folder"}], [{"id": "remote"}]])
            with patch("workbook.drive_setup.PROJECT_ROOT", root), \
                 patch("workbook.drive_setup.config_path", return_value=manifest):
                with self.assertRaisesRegex(DriveStorageError, "different data"):
                    migrate(client)
            self.assertFalse(manifest.exists())
            self.assertEqual(client.uploads, [])


if __name__ == "__main__":
    unittest.main()
